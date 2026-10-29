#!/usr/bin/env python3
"""Build the Sovereign Platinum reference workbook (items, wild Pokemon, trainers, statics, Totems, NPCs...)."""
import csv, json, re, sys, os, yaml, importlib.util
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

ROOT = "/home/user/pokemon-imperial-platinum"
OUT = sys.argv[1]
PK = "https://raw.githubusercontent.com/PokeAPI/sprites/master/sprites/pokemon/{}.png"
IT = "https://raw.githubusercontent.com/PokeAPI/sprites/master/sprites/items/{}.png"
TR = "https://play.pokemonshowdown.com/sprites/trainers/{}.png"

# ---- name -> PokeAPI id
species = {}
for r in csv.DictReader(open(f"{ROOT}/data/species.csv")):
    species[r["name"].lower()] = int(r["dex"])
forms = {}
for r in csv.DictReader(open(f"{ROOT}/data/forms.csv")):
    forms[r["identifier"]] = (int(r["form_id"]), r["form"])
    forms[r["form"].lower()] = (int(r["form_id"]), r["form"])
for r in csv.DictReader(open(f"{ROOT}/data/megas.csv")):
    forms[r["form"].lower()] = (int(r["form_id"]), r["form"])

def norm(n):
    return n.lower().replace("’", "'").strip()

def mon_id(name):
    n = norm(name)
    if n in species: return species[n]
    if n in forms: return forms[n][0]
    n2 = re.sub(r"\s*\(.*\)", "", n)
    if n2 in species: return species[n2]
    for k in (n.replace("-", " "), n.split()[-1] if n.split() else n):
        if k in species: return species[k]
    return None

def mon_img(name):
    i = mon_id(name)
    return PK.format(i) if i else ""

def item_img(name):
    slug = re.sub(r"[^a-z0-9 -]", "", norm(name).replace("é", "e")).replace(" ", "-")
    return IT.format(slug)

TRAINER_SPRITES = {"jupiter": "jupiter", "mars": "mars", "saturn": "saturn", "cyrus": "cyrus", "barry": "barry",
    "aaron": "aaron", "bertha": "bertha", "flint": "flint", "lucian": "lucian", "cynthia": "cynthia",
    "roark": "roark", "gardenia": "gardenia", "fantina": "fantina", "maylene": "maylene", "wake": "crasherwake",
    "byron": "byron", "candice": "candice", "volkner": "volkner", "dahlia": "dahlia"}
CLASS_SPRITES = {"EXPERT_M": "expert", "EXPERT_F": "expert-f", "ACE_TRAINER_M": "acetrainer", "ACE_TRAINER_F": "acetrainerf",
    "ACE_TRAINER_SNOW_M": "acetrainersnow", "ACE_TRAINER_SNOW_F": "acetrainersnowf", "GENTLEMAN": "gentleman",
    "IDOL": "idol", "LASS": "lass", "POKEFAN_F": "pokefanf", "POKEFAN_M": "pokefan", "RUIN_MANIAC": "ruinmaniac",
    "SCIENTIST_F": "scientistf", "SCIENTIST_M": "scientist", "HIKER": "hiker", "NINJA_BOY": "ninjaboy",
    "POLICEMAN": "policeman", "BUG_CATCHER": "bugcatcher", "BEAUTY": "beauty", "PSYCHIC": "psychic",
    "SOCIALITE": "socialite", "BLACK_BELT": "blackbelt", "GUITARIST": "guitarist", "DAHLIA": "dahlia",
    "GRUNT_M": "galacticgrunt", "MYSTERY_GIFT_DELIVERYMAN": "deliveryman", "ROUGHNECK": "roughneck"}

def trainer_img(name):
    for k, v in TRAINER_SPRITES.items():
        if k in name.lower(): return TR.format(v)
    return ""

def mapname(m):
    s = m.replace("_", " ").title()
    s = re.sub(r"\b(\d+)f\b", lambda x: x.group(1) + "F", s, flags=re.I)
    s = re.sub(r"\bB(\d)f\b", r"B\1F", s)
    s = s.replace("Hq", "HQ").replace(" Tv ", " TV ").replace("Trainers School", "Trainers' School").replace(" Gym Room 3", " Gym (Volkner's room)")
    return s.replace("Pokecenter", "Pokémon Center").replace("Pokemon", "Pokémon").replace("Mt ", "Mt. ")

wb = Workbook()
HEAD = Font(bold=True, color="FFFFFF")
from openpyxl.styles import PatternFill
FILL = PatternFill("solid", fgColor="1F3864")

def sheet(title, cols, rows, widths):
    ws = wb.create_sheet(title)
    hdr = ["Image"] + cols
    ws.append(hdr)
    for c in ws[1]:
        c.font = HEAD; c.fill = FILL
    for r in rows:
        url = r[0]
        ws.append([f'IMAGE("{url}")' if url else None] + [v if v != "" else None for v in r[1:]])
    ws.freeze_panes = "B2"
    from openpyxl.utils import get_column_letter
    ws.column_dimensions["A"].width = 12
    for i, w in enumerate(widths, 2):
        ws.column_dimensions[get_column_letter(i)].width = w
    if len(rows) < 1000:
        for row in ws.iter_rows(min_row=2):
            for c in row[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws

# ---------------- Items
items = []
txt = open(f"{ROOT}/docs/ItemChanges_Sovereign.md").read()
def md_tables(text, header_start):
    out = []
    for block in re.split(r"\n\n", text):
        lines = [l for l in block.split("\n") if l.startswith("|")]
        if len(lines) > 2 and lines[0].startswith(header_start):
            hdr = [h.strip() for h in lines[0].strip("|").split("|")]
            for l in lines[2:]:
                out.append(dict(zip(hdr, [c.strip().replace("**", "").replace("`", "") for c in l.strip("|").split("|")])))
    return out
for r in md_tables(txt, "| Item | Where | Purpose"):
    for nm in re.split(r",\s*|\s+and\s+", re.sub(r"\s*\(key items?\)", "", r["Item"])):
        cat = "Key item" if "key item" in r["Item"].lower() or r["Item"] in ("Mega Ring", "Light Stone, Dark Stone", "Prison Bottle", "Scroll of Darkness, Scroll of Waters", "Infinite Repel") else "Form item"
        items.append((item_img(nm), nm, cat, r["Where"], "", r["Purpose"]))
# crown pieces
for r in csv.DictReader(open(f"{ROOT}/data/item_sprites.csv")):
    if "Crown" in r["kind"] or "crown" in r["kind"].lower():
        items.append(("", r["item"], "Key item (" + r["kind"] + ")", r["source"], "Win the Totem trial" if "Totem" in r["source"] else "", "Part of the Sovereign Crown (Story.md)"))
y = yaml.safe_load(open(f"{ROOT}/data/items_sovereign.yaml"))
desc = {c["name"]: c["description"] for c in y["custom_items"]}
items = [(a, b, c, d, e, (f + " " + desc.get(b, "")).strip()) for a, b, c, d, e, f in items]
for nm in ("Crown", "Sovereign Crown", "Crown Diamond"):
    if not any(i[1] == nm for i in items):
        items.append(("", nm, "Key item", "Story (Story.md)", "", desc.get(nm, "")))
for nm in y["held_items"]:
    items.append((item_img(nm), nm, "Held item (Gen V–IX)", "Veilstone Dept. Store 2F, vitamin clerk → Held items", "", "Booster Energy ₽10,000" if nm == "Booster Energy" else ""))
for nm in y["gems"]:
    items.append((item_img(nm), nm, "Gem", "Veilstone Dept. Store 2F, vitamin clerk → Gems", "", "Boosts the first move of its type once"))
for nm in y["evolution_items"]:
    items.append((item_img(nm), nm, "Evolution item", "Snowpoint City evolution-item seller (₽10,000)" if not nm.startswith("Scroll") else "Legend Archivist, once Kubfu is caught", "", "See EvolutionChanges_Sovereign.md"))
for k, v in y["training"].items():
    for nm in v:
        items.append((item_img(nm), nm, "Mint" if k == "mints" else "Training item", "Route 208 Rare Berry seller (₽10,000)" if k == "mints" else "", "", "Changes the stat effect of the Nature" if k == "mints" else ""))
for k, v in y["form_items"].items():
    for nm in v:
        if not any(i[1] == nm for i in items):
            items.append((item_img(nm), nm, "Form item (" + k + ")", "", "", ""))
mega = md_tables(open(f"{ROOT}/docs/MegaEvolution.md").read(), "| Form | Stone")
for r in mega:
    items.append((item_img(r["Stone"]), r["Stone"], "Mega Stone", r["Where to get the stone"], "Mega Ring (Fantina, Relic Badge)", f"{r['Form']}: {r['Types']}, {r['Ability']}, BST {r['BST']}"))
for r in md_tables(txt, "| Shop | Where"):
    items.append(("", r["Shop"], "Shop", r["Where"], r["Opens"], r["Sells"]))
sheet("Items", ["Item", "Category", "Location", "Unlock / trigger", "Details"], items, [24, 20, 40, 26, 60])

# ---------------- Wild Pokemon
enc = json.load(open(f"{ROOT}/data/encounters_sovereign.json"))
G, W = enc["grass_rates"], enc["water_rates"]
wild = []
for a in enc["areas"]:
    lv = a["levels"].replace("Levels: ", "")
    if a["grass"]:
        for t in ("morning", "day", "night"):
            pass
        rates = {}
        for slot, s in enumerate(a["grass"]):
            for t in ("morning", "day", "night"):
                rates.setdefault(s[t], {"morning": 0, "day": 0, "night": 0})[t] += G[slot]
        for sp, r in rates.items():
            rs = f"{r['day']}%" if r["morning"] == r["day"] == r["night"] else f"Morning {r['morning']}% / Day {r['day']}% / Night {r['night']}%"
            wild.append((mon_img(sp), sp, a["name"], "Grass / cave", rs, lv, ""))
    for meth, slots in (a["water"] or {}).items():
        rates = {}
        for slot, sp in enumerate(slots):
            rates[sp] = rates.get(sp, 0) + W[slot]
        for sp, r in rates.items():
            wild.append((mon_img(sp), sp, a["name"], meth, f"{r}%", lv, ""))
    for sp, lvl in a["radar"]:
        wild.append((mon_img(sp), sp, a["name"], "Poké Radar", "", str(lvl), ""))
    for h in a["honey"]:
        sp = h[0] if isinstance(h, list) else h
        wild.append((mon_img(sp), sp, a["name"], "Honey tree", "", "", ""))
rare = yaml.safe_load(open(f"{ROOT}/data/sovereign_rare_wild.yaml"))
for r in rare["encounters"]:
    sp = r["species"].title()
    maps = sorted({mapname(m) for m in r["maps"]})
    loc = "Turnback Cave (pillar rooms)" if len(maps) > 3 else ", ".join(maps)
    lvl = r["level"] if not isinstance(r["level"], list) else f"{r['level'][0]}–{r['level'][1]}"
    wild.append((mon_img(sp), sp, loc, "Rare roll (" + r["terrain"] + ")", f"{r['rate']}%", str(lvl), "Only once " + rare["gates"][r["gate"]]["text"]))
park = yaml.safe_load(open(f"{ROOT}/data/regional_forms_park.yaml"))
for z in park["zones"]:
    for terr in ("grass", "water"):
        ents = z.get(terr) or []
        tot = sum(e[1] for e in ents)
        for ident, w, lo, hi in ents:
            fid, fname = forms.get(ident, (None, ident))
            wild.append((PK.format(fid) if fid else "", fname, f"Regional Forms Park — {z['name']} zone", "Tall grass" if terr == "grass" else "Surf / fishing", f"{100*w/tot:.0f}%", f"{lo}" if lo == hi else f"{lo}–{hi}", "After the Hall of Fame"))
merged = {}
for img, sp, loc, meth, rate, lv, trig in wild:
    if loc.startswith("Regional Forms Park") or meth.startswith("Rare roll"):
        merged[(sp, loc)] = [img, sp, f"{loc}: {meth} {rate}, Lv. {lv}", trig]
        continue
    m = merged.setdefault(sp, [img, sp, {}, ""])
    m[2].setdefault(loc, []).append(f"{meth} {rate}".strip() + (f" (Lv. {lv.split(' (')[0]})" if lv else ""))
wild = []
for m in merged.values():
    where = m[2] if isinstance(m[2], str) else "\n".join(f"{loc}: {', '.join(v)}" for loc, v in m[2].items())
    wild.append((m[0], m[1], where, m[3]))
wild.sort(key=lambda r: (mon_id(r[1]) or 99999, r[1]))
sheet("Wild Pokémon", ["Pokémon", "Where (method, rate, levels)", "Unlock / trigger"], wild, [18, 90, 30])

# ---------------- Trainers
tr = yaml.safe_load(open(f"{ROOT}/data/trainers_sovereign.yaml"))
def team(t):
    return "; ".join(f"{m['species']} Lv.{m['level']}" + (f" @ {m['item']}" if m.get("item") and m["item"] != "None" else "") for m in t)
trows = []
for x in tr["story"]:
    trows.append((trainer_img(x["trainer"]), x["trainer"], x["location"], x.get("format", "").title(), "Story", team(x.get("team") or []), x.get("note", "")))
for x in tr["league"]:
    rnd = "First Champion run" if x["round"] == 1 else "Rematch (after the Hall of Fame)"
    for fmt in ("singles", "doubles"):
        if x.get(fmt):
            trows.append((trainer_img(x["trainer"]), x["trainer"], "Pokémon League", fmt.title(), rnd, team(x[fmt]), x.get("note", "")))
sheet("Trainers", ["Trainer", "Location", "Format", "Unlock / trigger", "Team", "Details"], trows, [26, 30, 10, 22, 80, 50])

# ---------------- Static encounters
spec = importlib.util.spec_from_file_location("mx", os.path.join(os.path.dirname(__file__), "make_xlsx.py"))
src = open(spec.origin).read()
R = eval(src[src.index("R=[") + 2: src.index("\n]\n", src.index("R=[")) + 2], {"C": "Chronicle", "V": "Vanilla", "RP": "Renegade (ported)", "RN": "Renegade (not ported yet)", "T": "Totem trial"})
srows = []
for p, cat, grp, mp, spot, lvl, when, conf, notes in R:
    if cat == "Totem trial": continue
    first = p.split(" / ")[0]
    srows.append((mon_img(first), p, cat + ("" if grp in ("Vanilla", "Renegade") else f" — {grp}"), f"{mp} {spot}".strip(), when, f"Lv. {lvl}", notes))
sheet("Static encounters", ["Pokémon", "Category", "Location", "Unlock / trigger", "Level", "Details"], srows, [24, 26, 40, 34, 8, 60])

# ---------------- Totem trials
tt = yaml.safe_load(open(f"{ROOT}/data/totem_trials.yaml"))
TOTEM_IDS = {"Gumshoos": 10121, "Ribombee": 10150, "Vikavolt": 10122, "Mimikyu": 10144, "Marowak (Alolan)": 10149,
             "Raticate (Alolan)": 10093, "Lurantis": 10128, "Araquanid": 10153, "Salazzle": 10129, "Kommo-o": 10146, "Togedemaru": 10154}
crown = {r["source"].split("(")[0].replace("Totem ", "").split(" trial")[0].strip(): r["item"] for r in csv.DictReader(open(f"{ROOT}/data/item_sprites.csv"))}
ttrows = []
for t in tt["trials"]:
    reward = next((v for k, v in crown.items() if k.split()[0] in t["totem"]), "")
    ttrows.append((PK.format(TOTEM_IDS[t["totem"]]) if t["totem"] in TOTEM_IDS else mon_img(t["totem"]), "Totem " + t["totem"], t["location"] + " — " + t["map"],
                   f"After {t['after']} (before {t['before']})", f"Lv. {t['level']}",
                   f"Aura: {t['aura']}. Allies: {', '.join(t['helpers'])}. Reward: {reward}. {t['why']}"))
sheet("Totem trials", ["Totem", "Location", "Unlock / trigger", "Level", "Details"], ttrows, [24, 50, 40, 8, 70])

# ---------------- Gift Pokemon
gtxt = open(f"{ROOT}/docs/GiftPokemon_Sovereign.md").read()
grows = []
for r in md_tables(gtxt, "| Region | Pokémon"):
    for p in r["Pokémon"].split(", "):
        grows.append((mon_img(p), p, f"{r['Pokémon Center']} Pokémon Center, the Interviewer", r["Available"], f"Lv. {r['Level']}", f"{r['Region'].replace('*', '')} starter; quiz and battle, holds a type-boosting item"))
grows.append((PK.format(10116), "Battle Bond Greninja", "Fight Area Pokémon Center, Ace Trainer", "Post-game, once", "Lv. 50", "Water Shuriken, Aerial Ace, Double Team, Night Slash"))
for r in md_tables(gtxt, "| Pokémon | Lv. | Signature"):
    grows.append((mon_img(r["Pokémon"].split(",")[0].split(" ")[-1]), r["Pokémon"], "Fight Area Pokémon Center, event-gift deliveryman", "Post-game, each once", f"Lv. {r['Lv.']}", "Signature move: " + r["Signature move"]))
sheet("Gift Pokémon", ["Pokémon", "Location", "Unlock / trigger", "Level", "Details"], grows, [30, 44, 26, 8, 60])

# ---------------- Notable NPCs
nrows = []
for line in open("/tmp/claude-0/npcs.txt"):
    m, oid, gfx, xy, cm = line.rstrip("\n").split("|", 4)
    if "SIGN_" in oid or xy == "999,999" or oid.endswith("COUNTERPART") or gfx in ("MOSS_ROCK", "POKEBALL", "CAVE_PAINTING_SHARDS_LEFT") or "MEW_" in oid and gfx not in ("PACHIRISU",):
        if not ("MEW_" in oid): continue
    name = oid.replace("FIGHT_AREA_", "").replace("VEILSTONE_", "").replace("SOVEREIGN_", "").replace("_", " ").title()
    cm = re.sub(r"^Sovereign \([^)]*\):\s*", "", cm)
    img = TR.format(CLASS_SPRITES[gfx]) if gfx in CLASS_SPRITES else mon_img(gfx.title()) if mon_id(gfx.title()) else ""
    trig = ""
    for key, val in (("Chronicle", "Post-game (Chronicles)"), ("Hall of Fame", "Post-game"), ("eight badges", "All eight badges"), ("Type Master", "One Pokémon of the type, alone in the party"), ("Battle Zone", "Post-game (Battle Zone)"), ("Jirachi", "Jirachi event"), ("Mew event", "Mew event")):
        if key in cm: trig = val; break
    nrows.append((img, name, f"{mapname(m)} ({xy})", trig, cm or "—"))
extra = [
    ("expert", "Legend Archivist", "Canalave Library 2F", "", ""),
    ("pokefan", "Celestic Pokéfan (masks)", "Celestic Town, northeast house", "Ogerpon at the front of the party", "Gives the Wellspring, Hearthflame and Cornerstone Masks"),
    ("gentleman", "Mega Stone Collector", "Hearthome City, Bebe's house", "Relic Badge", "Sells 29 Mega Stones"),
    ("gentleman", "Mega Stone Trader", "Pastoria City, northeast house", "Fen Badge", "Sells 30 Mega Stones"),
    ("gentleman", "Mega Stone Sage", "Snowpoint City, west house", "Icicle Badge", "Sells 8 Mega Stones"),
    ("gentleman", "Mega Stone Curator", "Survival Area, south house", "Champion", "Sells 10 legendary/signature Mega Stones (₽30,000)"),
    ("fantina", "Leader Fantina (Mega Ring)", "Hearthome Gym", "Relic Badge", "Gives the Mega Ring"),
    ("reporter", "Interviewers (starters)", "Pokémon Centers: Jubilife, Oreburgh, Floaroma, Eterna, Hearthome, Veilstone, Pastoria, Canalave", "Badges (see Gift Pokémon)", "Quiz and battle for each region's starters"),
    ("oak", "Professor Oak (legendary birds)", "Eterna City, south house", "Post-game", "Releases Articuno, Zapdos and Moltres as roamers"),
    ("youngster", "Survival Area Youngster", "Survival Area", "Champion", "Gives 3 Nuggets once a day"),
]
names = {r[1] for r in nrows}
for s, n, loc, trig, det in extra:
    if n.split(" (")[0] not in names:
        nrows.append((TR.format(s), n, loc, trig, det))
sheet("Notable NPCs", ["NPC", "Location", "Unlock / trigger", "Details"], nrows, [28, 40, 30, 90])

# ---------------- TMs and HMs
txt = open(f"{ROOT}/RenegadePlatinumV130/Documentation/ItemChanges.txt", encoding="utf-8", errors="replace").read()
sec = txt[txt.index("TM Locations"):txt.index("Plate Locations")]
ren = {m[0]: (m[1].strip(), m[2].strip(), m[3].strip()) for m in re.findall(r"^((?:TM|HM)\d\d) (.+?)\s{2,}(.+?)\s{2,}(.+?)(?:\s+\*)?\s*$", sec, re.M)}
mtype = {}
TYPES = {1: "normal", 2: "fighting", 3: "flying", 4: "poison", 5: "ground", 6: "rock", 7: "bug", 8: "ghost", 9: "steel", 10: "fire", 11: "water", 12: "grass", 13: "electric", 14: "psychic", 15: "ice", 16: "dragon", 17: "dark", 18: "fairy"}
for r in csv.DictReader(open("/home/user/pokeapi/pokeapi/data/v2/csv/moves.csv")):
    mtype[r["identifier"]] = TYPES.get(int(r["type_id"]), "normal")
src = {}
for line in open("/tmp/claude-0/tm_sources.txt"):
    item, kind, m, where, flag = line.rstrip("\n").split("|")
    if m.startswith("unused") or m.endswith("_low_water") or m.endswith("_drained") or "_dp_" in m or m.startswith("ruin_maniac_cave"):
        continue
    src.setdefault(item[5:], []).append((kind, m, where))
NOTES = {
    "TM10": "For beating both students in the Trainers' School",
    "TM13": "An Ace Trainer in the Route 216 lodge, with five badges",
    "TM15": "An Ace Trainer in Sunyshore City's east house",
    "TM17": "The Cowgirl in the Great Marsh gate, for five catches in one Safari Game",
    "TM27": "Cheryl, at the end of Eterna Forest",
    "TM29": "Lucian, who comes down from the Solaceon Ruins once",
    "TM42": "An idol in the Jubilife TV lobby (the Survival Area house gives one too)",
    "TM44": "An old woman in Hearthome's foreign building",
    "TM52": "Maylene, in her Route 217 scene",
    "TM54": "The aide in the Cycling Road north gate, for 150 Pokémon seen",
    "TM55": "A cook in the Poffin House, for a Burn Heal",
    "TM26": "Bertha, in the Veilstone Department Store 5F café",
    "TM28": "A Ninja Boy in the Underground Man's house",
    "TM30": "Fantina, after her Gym battle",
    "TM40": "An old swordsman in Floaroma Town's middle house",
    "TM50": "Flint, in Sunyshore City before he leaves for the League",
    "TM80": "The Underground Man's house, for showing an Aerodactyl",
    "TM81": "Aaron, after the Route 208 battle",
    "TM85": "Mira, at the exit of Wayward Cave, after her quest",
    "HM04": "Riley, after the Iron Island tag battle",
    "HM07": "Wake, after his Gym battle",
    "HM08": "Route 217 item ball for now (Renegade has Candice give it after a Snowpoint Temple scene; that scene waits on the owner)",
}
tms = []
for tm in sorted(ren, key=lambda t: (t[:2] == "HM", t)):
    move, rloc, robt = ren[tm]
    ty = mtype.get(re.sub(r"[^a-z0-9 -]", "", move.lower()).replace(" ", "-"), "normal")
    places = []
    for kind, m, where in src.get(tm, []):
        places.append(f"{mapname(m)} ({where.replace(',', ', ')})" if kind == "ball" else mapname(m))
    places = list(dict.fromkeys(places))
    method = "; ".join(dict.fromkeys("Item ball" if k == "ball" else "Gift" for k, _, _ in src.get(tm, []))) or robt
    note = NOTES.get(tm, robt if robt != "Item on ground" and method == "Gift" else "")
    tms.append((IT.format(f"tm-{ty}"), f"{tm} {move}", ty.title(), "; ".join(places) or rloc, method, note))
sheet("TMs and HMs", ["TM / HM", "Type", "Location", "Method", "Unlock / details"], tms, [26, 10, 46, 14, 70])

# ---------------- README
ws = wb["Sheet"]; ws.title = "Read me"
for line in [
    ["Pokémon Sovereign Platinum — reference sheets"],
    ["Built from the repository (data/ and docs/) on 2026-10-24. Tabs: Items, Wild Pokémon, Trainers, Static encounters, Totem trials, Gift Pokémon, Notable NPCs, TMs and HMs."],
    [""],
    ["Images: Google's importer won't create image formulas, so column A of each tab holds the formula as text: IMAGE(\"link\")."],
    ["To show every image at once: Edit > Find and replace (Ctrl+H). Find:  IMAGE(    Replace with:  =IMAGE(    Search: All sheets. Replace all."],
    ["Some items and NPCs have no public image, so their cell stays blank."],
    [""],
    ["Wild Pokémon rates are the share of encounters for that method; time-of-day rates are shown when they differ."],
    ["Trainers lists the trainers whose teams Sovereign changed; everyone else follows Renegade Platinum's documentation."],
]:
    ws.append(line)
ws["A1"].font = Font(bold=True, size=14)
ws.column_dimensions["A"].width = 140
wb.save(OUT)
print({s.title: s.max_row - 1 for s in wb.worksheets})
