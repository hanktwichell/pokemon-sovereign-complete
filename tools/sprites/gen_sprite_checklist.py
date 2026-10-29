#!/usr/bin/env python3
"""Write docs/SpriteChecklist.xlsx: which sprites we have for all 1025 species (and every Mega).

A check mark means we have the sprite; a highlighted empty cell means we still need it.
  * #1-493 already exist in Renegade Platinum (the owner's rule: keep Renegade's sprites).
  * #494-1025: every pack under assets/sprites/packs/ is merged:
      - DPPt Style Gen VI and Beyond (Essentials layout, backs) via data/sprite_index.csv
      - any "dex folder" pack: NNNN-Name/{front,back,shiny_front,shiny_back}.png
  * Box/party icons for added species: placeholders converted from PokeAPI icons
    (tools/sprites/make_placeholder_icons.py, data/icon_placeholders.csv).
  * New items' bag icons: data/item_sprites.csv (have_icon = 1 once the icon is in).
The Notes column says whether each sprite is DS-ready (80x80, <=15 colours), can be cropped
losslessly to 80x80, or needs a resize/redraw.
Re-run after adding sprites (and after re-running tools/sprites/import_sprite_pack.py).

Usage: python3 tools/sprites/gen_sprite_checklist.py
"""
import csv
import os

from PIL import Image
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CHECK = "✓"
FONT = "Arial"
MISSING = PatternFill("solid", start_color="F4CCCC")
HEADER = PatternFill("solid", start_color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SPRITE_COLS = ["Non-shiny front sprite", "Non-shiny back sprite", "Shiny front sprite", "Shiny back sprite",
               "Box/party screen sprite"]
# Owner-supplied Mega fronts (assets/sprites/owner_supplied/).
OWNER_MEGA_FRONTS = {"Mega Garchomp Z": "Owner-supplied front + shiny front; 116×95 art must be redrawn to fit the DS 80×80 frame."}


PACKS = os.path.join(ROOT, "assets", "sprites", "packs")
PACK_LABELS = {"gen5_9_part1": "Gen5-9 Sprites part 1", "gen5_9_part2": "Gen5-9 Sprites part 2"}
KINDS = {"front": "front", "back": "back", "shiny_front": "front_shiny", "shiny_back": "back_shiny"}


def ds_status(path):
    """'ready' | 'crop' | 'resize', plus a short reason."""
    im = Image.open(path).convert("RGBA")
    w, h = im.size
    colours = len({p for p in im.get_flattened_data() if p[3] > 0})
    bbox = im.getbbox() or (0, 0, 0, 0)
    bw, bh = bbox[2] - bbox[0], bbox[3] - bbox[1]
    if colours > 15:
        return "resize", f"{colours} colours"
    if (w, h) == (80, 80):
        return "ready", ""
    if bw <= 80 and bh <= 80:
        return "crop", f"{w}x{h} canvas, art fits 80x80"
    return "resize", f"{w}x{h}, art {bw}x{bh}"


def dex_folder_packs():
    """{(dex, kind): [(pack_label, path, status, reason)]} for packs in NNNN-Name/ layout."""
    found = {}
    if not os.path.isdir(PACKS):
        return found
    for pack in sorted(os.listdir(PACKS)):
        root = os.path.join(PACKS, pack)
        if not os.path.isdir(root):
            continue
        for d in sorted(os.listdir(root)):
            m = d.split("-", 1)
            if len(m) != 2 or not m[0].isdigit():
                continue
            for fn, kind in KINDS.items():
                p = os.path.join(root, d, fn + ".png")
                if os.path.exists(p):
                    st, why = ds_status(p)
                    found.setdefault((int(m[0]), kind), []).append((PACK_LABELS.get(pack, pack), p, st, why))
    return found


def load():
    species = list(csv.DictReader(open(os.path.join(ROOT, "data", "species.csv"), encoding="utf-8")))
    megas = list(csv.DictReader(open(os.path.join(ROOT, "data", "megas.csv"), encoding="utf-8")))
    index = list(csv.DictReader(open(os.path.join(ROOT, "data", "sprite_index.csv"), encoding="utf-8")))
    files = {}
    for e in index:
        if e["category"] in ("base", "mega") and "(female)" not in e["form"]:
            files[(e["form"], e["kind"])] = e
    return species, megas, files


def pack_note(entries):
    notes = []
    for label, e in entries:
        if e and e["ds_ready"] != "1":
            why = f"{e['colours']} colours" if int(e["colours"]) > 15 else "not a clean 2× upscale"
            notes.append(f"{label}: {why}")
    return ("Needs DS conversion — " + "; ".join(notes)) if notes else ""


def style_header(ws, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"


def write_row(ws, r, values, sprite_start, n_sprites):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT)
        cell.border = BORDER
        if sprite_start <= c < sprite_start + n_sprites:
            cell.alignment = Alignment(horizontal="center")
            if v != CHECK:
                cell.fill = MISSING


def main():
    species, megas, files = load()
    folder = dex_folder_packs()
    icons = {}
    icon_csv = os.path.join(ROOT, "data", "icon_placeholders.csv")
    if os.path.exists(icon_csv):
        icons = {int(r["dex"]): r for r in csv.DictReader(open(icon_csv, encoding="utf-8"))}
    wb = Workbook()

    # ---- Sheet 1: all 1025 species
    ws = wb.active
    ws.title = "Sprite Checklist"
    headers = ["Dex #", "Pokémon", "Generation", *SPRITE_COLS, "Source", "Notes"]
    style_header(ws, headers)
    for i, s in enumerate(species, start=2):
        if s["in_renegade"] == "1":
            marks = [CHECK] * 5
            source = "Renegade Platinum ROM (existing sprites)"
            note = ""
        else:
            dex = int(s["dex"])
            ess = {"back": files.get((s["name"], "back")), "back_shiny": files.get((s["name"], "back_shiny"))}
            marks, sources, notes = [], [], []
            for col, kind in zip(SPRITE_COLS[:4], ("front", "back", "front_shiny", "back_shiny")):
                opts = list(folder.get((dex, kind), []))
                if ess.get(kind):
                    e = ess[kind]
                    st = "ready" if e["ds_ready"] == "1" else "resize"
                    why = "" if st == "ready" else (f"{e['colours']} colours" if int(e["colours"]) > 15 else "not a clean 2x upscale")
                    opts.append(("DPPt Style Gen VI and Beyond", None, st, why))
                marks.append(CHECK if opts else "")
                if opts:
                    rank = {"ready": 0, "crop": 1, "resize": 2}
                    best = min(opts, key=lambda o: rank[o[2]])
                    for o in opts:
                        if o[0] not in sources:
                            sources.append(o[0])
                    if best[2] != "ready":
                        label = kind.replace("_shiny", "").replace("front", "front").replace("back", "back")
                        label = ("shiny " if "shiny" in kind else "") + label
                        notes.append(f"{label}: {'crop to 80x80' if best[2] == 'crop' else 'resize/redraw'} ({best[3]})")
            icon = icons.get(dex)
            marks.append(CHECK if icon else "")  # box/party icon
            if icon:
                sources.append(icon["source"] + " (placeholder icon)")
            source = "; ".join(sources)
            note = "; ".join(notes)
        write_row(ws, i, [int(s["dex"]), s["name"], int(s["generation"]), *marks, source, note], 4, 5)
    last = len(species) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    for col, w in zip("ABCDEFGHIJ", [7, 18, 11, 14, 14, 14, 14, 14, 38, 50]):
        ws.column_dimensions[col].width = w

    # ---- Sheet 2: Mega Evolutions / Primal Reversions
    wm = wb.create_sheet("Mega & Primal Forms")
    mheaders = ["Dex #", "Form", "Non-shiny front sprite", "Non-shiny back sprite", "Shiny front sprite",
                "Shiny back sprite", "Source", "Notes"]
    style_header(wm, mheaders)
    wm.cell(row=1, column=2).comment = Comment(
        "Mega/Primal forms have no box/party column: the party and PC show the base form's icon.", "Sovereign Platinum")
    for i, m in enumerate(megas, start=2):
        back = files.get((m["form"], "back"))
        sback = files.get((m["form"], "back_shiny"))
        front = m["form"] in OWNER_MEGA_FRONTS
        marks = [CHECK if front else "", CHECK if back else "", CHECK if front else "", CHECK if sback else ""]
        srcs = []
        if back or sback:
            srcs.append("DPPt Style Gen VI and Beyond pack (backs)")
        if front:
            srcs.append("Owner-supplied (fronts)")
        note = "; ".join(x for x in (OWNER_MEGA_FRONTS.get(m["form"], ""), pack_note([("back", back), ("shiny back", sback)])) if x)
        write_row(wm, i, [int(m["dex"]), m["form"], *marks, "; ".join(srcs), note], 3, 4)
    mlast = len(megas) + 1
    wm.auto_filter.ref = f"A1:H{mlast}"
    for col, w in zip("ABCDEFGH", [7, 32, 14, 14, 14, 14, 44, 60]):
        wm.column_dimensions[col].width = w

    # ---- Sheet 3: bag icons for new items
    items = list(csv.DictReader(open(os.path.join(ROOT, "data", "item_sprites.csv"), encoding="utf-8")))
    wi = wb.create_sheet("Item Icons")
    style_header(wi, ["Item", "Kind", "Where it comes from", "Bag icon", "Notes"])
    wi.cell(row=1, column=4).comment = Comment("Bag icon: 32×32, 16 colours, as Platinum's item icons.", "Sovereign Platinum")
    for i, it in enumerate(items, start=2):
        write_row(wi, i, [it["item"], it["kind"], it["source"], CHECK if it["have_icon"] == "1" else "", it["notes"]], 4, 1)
    ilast = len(items) + 1
    for col, w in zip("ABCDE", [18, 16, 58, 12, 50]):
        wi.column_dimensions[col].width = w

    # ---- Sheet 4: summary (live formulas)
    su = wb.create_sheet("Summary", 0)
    su["A1"] = "Sovereign Platinum — Sprite Aggregation Summary"
    su["A1"].font = Font(name=FONT, bold=True, size=14)
    su["A2"] = "✓ = we have the sprite. Red-highlighted empty cell = still needed. Counts update as you add ✓ marks."
    su["A3"] = ("#1–493 are already in Renegade Platinum (we keep its sprites). Added species: DPPt Style Gen VI and "
                "Beyond pack (backs) + Gen5-9 Sprites part 1 (fronts/backs/some shinies). Box/party icons for added species are "
                "placeholders converted from PokeAPI's icons. "
                "Notes say which sprites still need cropping or resizing for the DS.")
    for r in (2, 3):
        su[f"A{r}"].font = Font(name=FONT, italic=True)
    hdr = ["Sprite", "All species: have", "All species: missing", "Added species (#494–1025): have",
           "Added species: missing", "Megas: have", "Megas: missing"]
    for c, h in enumerate(hdr, 1):
        cell = su.cell(row=5, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BORDER
    sc = "'Sprite Checklist'"
    mg = "'Mega & Primal Forms'"
    added_first = next(i for i, s in enumerate(species, start=2) if s["in_renegade"] == "0")
    mega_col = {SPRITE_COLS[0]: "C", SPRITE_COLS[1]: "D", SPRITE_COLS[2]: "E", SPRITE_COLS[3]: "F"}
    for r, name in enumerate(SPRITE_COLS, start=6):
        col = get_column_letter(4 + SPRITE_COLS.index(name))
        all_rng = f"{sc}!${col}$2:${col}${last}"
        add_rng = f"{sc}!${col}${added_first}:${col}${last}"
        row = [name,
               f'=COUNTIF({all_rng},"{CHECK}")', f'=ROWS({all_rng})-B{r}',
               f'=COUNTIF({add_rng},"{CHECK}")', f'=ROWS({add_rng})-D{r}']
        if name in mega_col:
            mrng = f"{mg}!${mega_col[name]}$2:${mega_col[name]}${mlast}"
            row += [f'=COUNTIF({mrng},"{CHECK}")', f'=ROWS({mrng})-F{r}']
        else:
            row += ["n/a", "n/a"]
        for c, v in enumerate(row, 1):
            cell = su.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, bold=(c == 1))
            cell.border = BORDER
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
    r = 6 + len(SPRITE_COLS)
    irng = f"'Item Icons'!$D$2:$D${ilast}"
    for c, v in enumerate(["New item bag icons", f'=COUNTIF({irng},"{CHECK}")', f'=ROWS({irng})-B{r}'], 1):
        cell = su.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, bold=(c == 1))
        cell.border = BORDER
        if c > 1:
            cell.alignment = Alignment(horizontal="center")
    su.column_dimensions["A"].width = 26
    for col in "BCDEFG":
        su.column_dimensions[col].width = 17
    su.row_dimensions[5].height = 32

    out = os.path.join(ROOT, "docs", "SpriteChecklist.xlsx")
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
